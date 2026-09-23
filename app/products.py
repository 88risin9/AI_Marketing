"""Product records and conservative, preview-first spreadsheet imports.

Identifiers and units remain strings exactly as supplied. Importing never writes
the database; the API validates the whole selected batch before committing it.
"""

from __future__ import annotations

import copy
import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 100
CURRENCIES = {"CNY", "USD", "EUR", "GBP", "JPY", "AUD", "CAD"}
ASSETS = Path(__file__).resolve().parent.parent / "assets"

# The first row is deliberately machine readable; instructions live separately.
COLUMNS = [
    ("供应商", "supplier"), ("产品名称", "name"), ("英文名称", "name_en"),
    ("型号", "model"), ("分类", "category"), ("极数", "specs.poles"),
    ("额定电流(A)", "specs.current_a"), ("额定电压(V)", "specs.voltage_v"),
    ("脱扣曲线", "specs.curve"), ("分断能力(kA)", "specs.breaking_ka"),
    ("计量单位", "unit"), ("采购价", "purchase_price"), ("采购币种", "currency"),
    ("最小起订量", "moq"), ("交期", "lead_time"), ("资料来源", "source"),
    ("资料更新时间", "updated_at"), ("扩展参数(JSON)", "specs"),
]
LABELS = {key: label for label, key in COLUMNS}
ALIASES = {label: key for label, key in COLUMNS}
ALIASES.update({key: key for _, key in COLUMNS})
ALIASES.update({
    "单位": "unit", "币种": "currency", "规格参数": "specs", "更新时间": "updated_at",
    "MOQ": "moq", "moq": "moq", "额定电流": "specs.current_a",
    "额定电压": "specs.voltage_v", "分断能力": "specs.breaking_ka",
    "产品大类": "specs.product_family", "资料类型": "specs.record_type",
    "目录摘要": "specs.source_specification", "待确认事项": "specs.review_notes",
    **{key: "specs." + key for key in ("poles", "current_a", "voltage_v", "curve", "breaking_ka")},
})

DEMO_PRODUCTS = [
    {
        "supplier": "虚构示例·海岚电器", "name": "虚构示例·小型断路器", "name_en": "Miniature circuit breaker",
        "model": "DEMO-MCB-C16-1P", "category": "MCB",
        "specs": {"poles": "1P", "current_a": "16", "voltage_v": "230", "curve": "C", "breaking_ka": "6"},
        "unit": "pcs", "purchase_price": "8.50", "currency": "CNY", "moq": "100",
        "lead_time": "15 days after order confirmation", "source": "虚构资料，仅供演示；无实际认证、库存或供货承诺", "updated_at": "2026-09-15",
    },
    {
        "supplier": "虚构示例·海岚电器", "name": "虚构示例·小型断路器", "name_en": "Miniature circuit breaker",
        "model": "DEMO-MCB-C32-2P", "category": "MCB",
        "specs": {"poles": "2P", "current_a": "32", "voltage_v": "400", "curve": "C", "breaking_ka": "6"},
        "unit": "pcs", "purchase_price": "19.20", "currency": "CNY", "moq": "100",
        "lead_time": "20 days after order confirmation", "source": "虚构资料，仅供演示；参数与第一个型号不同", "updated_at": "2026-09-15",
    },
    {
        "supplier": "虚构示例·江屿电器", "name": "虚构示例·资料待补小型断路器", "name_en": "Miniature circuit breaker",
        "model": "DEMO-MCB-UNKNOWN", "category": "MCB",
        "specs": {"poles": "1P", "current_a": "16", "voltage_v": "", "curve": "", "breaking_ka": ""},
        "unit": "pcs", "purchase_price": "", "currency": "", "moq": "",
        "lead_time": "", "source": "虚构资料，仅供演示；缺失内容必须向供应商确认", "updated_at": "2026-09-15",
    },
]


def example_products():
    """Fresh copies prevent demonstration edits mutating subsequent seeds."""
    return [validate_product(p) for p in copy.deepcopy(DEMO_PRODUCTS)]


def _text(value, label, preserve=False):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, bool)):
        raise ValueError(f"{label}必须填写文本")
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    result = str(value)
    if len(result) > 10000:
        raise ValueError(f"{label}内容过长（最多 10000 个字符）")
    if "\x00" in result:
        raise ValueError(f"{label}包含无效字符")
    return result if preserve else result.strip()


def _decimal(value, label, positive=False):
    raw = _text(value, label)
    if not raw:
        return ""
    if len(raw) > 30 or not re.fullmatch(r"\+?\d+(?:\.\d+)?", raw):
        raise ValueError(f"{label}须为普通十进制数，不含币种、千位逗号或单位")
    try:
        number = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"{label}不是有效数值") from exc
    if not number.is_finite() or (number <= 0 if positive else number < 0):
        raise ValueError(f"{label}须{'大于零' if positive else '大于或等于零'}")
    if number > Decimal("1000000000000") or -number.as_tuple().exponent > 8:
        raise ValueError(f"{label}数值过大或超过 8 位小数")
    return format(number, "f")


def validate_product(payload):
    """Return a safe canonical record, allowing omissions but rejecting bad data."""
    if not isinstance(payload, dict):
        raise ValueError("产品资料必须为对象")
    product = {}
    for field in ("supplier", "name", "name_en", "model", "category", "unit", "lead_time", "source", "updated_at"):
        if field in {"model", "unit"} and payload.get(field) is not None and not isinstance(payload[field], str):
            raise ValueError(f"{LABELS[field]}必须为文本，以保留前导零和原始写法")
        product[field] = _text(payload.get(field), LABELS.get(field, field), field in {"model", "unit"})
    product["purchase_price"] = _decimal(payload.get("purchase_price"), "采购价")
    product["moq"] = _decimal(payload.get("moq"), "最小起订量", positive=True)
    product["currency"] = _text(payload.get("currency"), "采购币种").upper()
    if product["currency"] and product["currency"] not in CURRENCIES:
        raise ValueError("采购币种仅支持 CNY、USD、EUR、GBP、JPY、AUD、CAD")
    specs = payload.get("specs")
    if specs is None or specs == "":
        specs = {}
    if isinstance(specs, str):
        try:
            specs = json.loads(specs)
        except (ValueError, TypeError) as exc:
            raise ValueError('扩展参数须为 JSON 对象，例如 {"frequency_hz":"50/60"}') from exc
    if not isinstance(specs, dict) or len(specs) > 100:
        raise ValueError("规格参数须为对象，最多 100 个参数")
    product["specs"] = {}
    for key, value in specs.items():
        if not isinstance(key, str) or not key.strip() or len(key) > 100:
            raise ValueError("规格参数名称须为 1 至 100 个字符")
        normalized = _text(value, f"参数 {key}")
        product["specs"][key] = normalized
    missing = [LABELS[key] for key in ("supplier", "name", "name_en", "model", "category", "unit", "purchase_price", "currency", "moq", "lead_time", "source", "updated_at") if not product[key].strip()]
    if product["category"].upper() == "MCB":
        for key in ("poles", "current_a", "voltage_v", "curve", "breaking_ka"):
            if not product["specs"].get(key, "").strip():
                missing.append(LABELS["specs." + key])
    product["missing"] = missing
    if "archived" in payload:
        if type(payload["archived"]) is not bool:
            raise ValueError("产品归档标记必须为布尔值")
        product["archived"] = payload["archived"]
    if "archived_at" in payload:
        archived_at = _text(payload["archived_at"], "产品归档时间")
        if archived_at:
            try:
                datetime.fromisoformat(archived_at)
            except ValueError as exc:
                raise ValueError("产品归档时间格式不正确") from exc
        product["archived_at"] = archived_at
    return product


def _header(value):
    name = str(value or "").strip().lstrip("\ufeff")
    if name in ALIASES:
        return ALIASES[name]
    if name.lower() in ALIASES:
        return ALIASES[name.lower()]
    for prefix in ("参数:", "参数：", "specs."):
        if name.startswith(prefix) and name[len(prefix):].strip():
            return "specs." + name[len(prefix):].strip()
    return None


def _load_csv(content):
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if text is None or "\x00" in text:
        raise ValueError("CSV 编码无法识别，请另存为 UTF-8 CSV")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t") if text.strip() else csv.excel
    except csv.Error:
        dialect = csv.excel
    result = []
    try:
        for values in csv.reader(io.StringIO(text), dialect=dialect, strict=True):
            if len(result) > MAX_ROWS or len(values) > MAX_COLUMNS:
                raise ValueError(f"最多导入 {MAX_ROWS} 行、{MAX_COLUMNS} 列")
            result.append([(value, None) for value in values])
    except csv.Error as exc:
        raise ValueError("CSV 格式不正确，请检查引号和分隔符") from exc
    return result, []


def _load_xlsx(content):
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Excel 解压后过大，请拆分资料后导入")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel 文件无法读取，请使用未加密的 .xlsx 文件") from exc
    try:
        sheet = wb["产品资料"] if "产品资料" in wb.sheetnames else wb.worksheets[0]
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ValueError(f"最多导入 {MAX_ROWS} 行、{MAX_COLUMNS} 列，请清除多余格式或拆分文件")
        result = []
        # Some valid writers omit <dimension>; cap streaming iteration ourselves.
        for row_index, row in enumerate(sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1)):
            values = [(cell.value, cell.data_type) for cell in row]
            while values and values[-1][0] in (None, ""):
                values.pop()
            if len(values) > MAX_COLUMNS or (row_index > MAX_ROWS and values):
                raise ValueError(f"最多导入 {MAX_ROWS} 行、{MAX_COLUMNS} 列")
            result.append(values)
        notes = []
        others = [s.title for s in wb.worksheets if s.title != sheet.title and s.title not in {"填写说明", "说明"}]
        if others:
            notes.append(f"本次只读取工作表“{sheet.title}”，其他工作表未导入：" + "、".join(others))
        return result, notes
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel 工作表内容损坏，请重新另存为 .xlsx 后重试") from exc
    finally:
        wb.close()


def parse_import(filename, content):
    """Preview only; row errors never partially import into persistent storage."""
    empty = {"rows": [], "errors": [], "warnings": []}
    if not isinstance(content, bytes) or not content:
        return {**empty, "errors": ["文件为空"]}
    if len(content) > MAX_BYTES:
        return {**empty, "errors": ["文件超过 5 MB，请拆分后导入"]}
    extension = Path(filename or "").suffix.lower()
    try:
        if extension == ".csv":
            data, notes = _load_csv(content)
        elif extension == ".xlsx":
            data, notes = _load_xlsx(content)
        else:
            raise ValueError("仅支持 .csv 和 .xlsx；旧版 .xls 请另存为 .xlsx")
    except (ValueError, csv.Error) as exc:
        return {**empty, "errors": [str(exc)]}
    if not data:
        return {**empty, "errors": ["文件中没有表头"]}
    # Excel's formatted trailing cells are not extra input columns.
    header_values = list(data[0])
    while header_values and header_values[-1][0] in (None, ""):
        header_values.pop()
    headers = [_header(value) for value, _ in header_values]
    header_errors = []
    if not any(headers):
        header_errors.append("第 1 行没有可识别表头，请使用下载的模板")
    seen = set()
    for index, key in enumerate(headers):
        label = header_values[index][0]
        if not key:
            header_errors.append(f"第 {index + 1} 列表头“{label or ''}”无法识别；扩展参数请用 参数:名称")
        elif key in seen:
            header_errors.append(f"表头“{label}”重复映射到 {key}")
        else:
            seen.add(key)
    if header_errors:
        return {**empty, "errors": header_errors, "warnings": notes}
    rows = []
    identities = set()
    model_identities = set()
    for row_number, values in enumerate(data[1:], start=2):
        if all(value in (None, "") for value, _ in values):
            continue
        row_errors, warnings = [], []
        product = {"specs": {}}
        if any(value not in (None, "") for value, _ in values[len(headers):]):
            row_errors.append("存在没有表头的数据列，请补齐表头")
        expanded = {}
        for col_index, key in enumerate(headers):
            value, cell_type = values[col_index] if col_index < len(values) else (None, None)
            if cell_type == "f":
                row_errors.append(f"{header_values[col_index][0]}包含公式，请粘贴为纯数值或文本")
            if key in {"model", "unit"} and value is not None and not isinstance(value, str):
                row_errors.append(f"{LABELS[key]}不是文本，Excel 可能已改变前导零或日期；请对照原资料修正为文本")
            if isinstance(value, str) and re.match(r"^[\s]*[=+@-]", value):
                warnings.append(f"{header_values[col_index][0]}以公式符号开头，将按原始文本保存，请确认")
            if isinstance(value, (date, datetime)):
                value = value.isoformat()
            if key == "specs":
                if value not in (None, ""):
                    try:
                        parsed = json.loads(str(value))
                        if not isinstance(parsed, dict):
                            raise ValueError()
                        expanded.update(parsed)
                    except (ValueError, TypeError):
                        row_errors.append('扩展参数须为 JSON 对象，例如 {"frequency_hz":"50/60"}')
            elif key.startswith("specs."):
                product["specs"][key[6:]] = value if value is not None else ""
            else:
                product[key] = value if value is not None else ""
        for key, value in expanded.items():
            if key in product["specs"] and product["specs"][key] not in (None, "") and str(product["specs"][key]) != str(value):
                row_errors.append(f"扩展参数 {key} 与独立参数列冲突")
            elif key not in product["specs"] or product["specs"][key] in (None, ""):
                product["specs"][key] = value
        try:
            product = validate_product(product)
            if product["missing"]:
                warnings.append("资料待补：" + "、".join(product["missing"]))
        except ValueError as exc:
            row_errors.append(str(exc))
        model_identity = (str(product.get("supplier", "")), str(product.get("model", "")))
        identity = model_identity + (json.dumps(product.get("specs", {}), ensure_ascii=False, sort_keys=True), str(product.get("unit", "")))
        if model_identity[1].strip() and identity in identities:
            warnings.append("文件中出现相同供应商、型号、规格和单位，确认是否为重复记录")
        elif model_identity[1].strip() and model_identity in model_identities:
            warnings.append("同型号存在不同规格或单位，作为独立变体保留，请逐行核对")
        identities.add(identity)
        model_identities.add(model_identity)
        rows.append({"row": row_number, "product": product, "errors": list(dict.fromkeys(row_errors)), "warnings": warnings})
    return {"rows": rows, "errors": [], "warnings": notes}


def csv_safe(value):
    """Protect spreadsheet readers when emitting CSV, without changing records."""
    text = "" if value is None else str(value)
    return "'" + text if re.match(r"^[\s]*[=+@-]", text) else text


def template_bytes(format, demo=False):
    if format.lower() not in {"csv", "xlsx"}:
        raise ValueError("模板格式仅支持 csv 或 xlsx")
    stem = "example" if demo else "products"
    return (ASSETS / f"{stem}.{format.lower()}").read_bytes()
